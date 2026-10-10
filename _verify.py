# -*- coding: utf-8 -*-
import docx
d = docx.Document(r'E:\2024级计科软工五班大三课程\Android开发\实验二\实验二+杨楠2024110539.docx')
out = []
for i, p in enumerate(d.paragraphs):
    t = p.text.strip()
    # 数图片
    imgs = len(p._element.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip'))
    if t or imgs:
        out.append(f"{i} | img={imgs} | {t[:90]}")
with open(r'E:\Android exeperise\exeperise2\_verify.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('paras:', len(d.paragraphs))
