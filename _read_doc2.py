# -*- coding: utf-8 -*-
import docx
src = r'E:\2024级计科软工五班大三课程\Android开发\实验二\实验二+UI控件代码生成以及MVC.docx'
d = docx.Document(src)
out = []
for i in range(25, min(62, len(d.paragraphs))):
    p = d.paragraphs[i]
    out.append(f"{i} | {repr(p.text[:80])}")
with open(r'E:\Android exeperise\exeperise2\_doc_struct.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('done', len(out))
