# -*- coding: utf-8 -*-
import docx
src = r'E:\2024级计科软工五班大三课程\Android开发\实验二\实验二+UI控件代码生成以及MVC.docx'
d = docx.Document(src)
out = []
out.append(f"=== 段落总数: {len(d.paragraphs)} ===")
for i, p in enumerate(d.paragraphs):
    out.append(f"[P{i}] style={p.style.name} | {repr(p.text)}")
out.append(f"\n=== 表格数: {len(d.tables)} ===")
for ti, tbl in enumerate(d.tables):
    out.append(f"--- Table {ti}: {len(tbl.rows)} rows x {len(tbl.columns)} cols ---")
    for ri, row in enumerate(tbl.rows):
        cells = [c.text.strip()[:50] for c in row.cells]
        out.append(f"  R{ri}: {cells}")
# 节/页眉
out.append(f"\n=== 节数: {len(d.sections)} ===")
for si, sec in enumerate(d.sections):
    out.append(f"Sec{si}: page {sec.page_width},{sec.page_height}")
with open(r'E:\Android exeperise\exeperise2\_full_struct.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('done')
