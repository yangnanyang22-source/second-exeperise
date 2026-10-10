# -*- coding: utf-8 -*-
import shutil, os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = r'E:\2024级计科软工五班大三课程\Android开发\实验二\实验二+UI控件代码生成以及MVC.docx'
dst = r'E:\2024级计科软工五班大三课程\Android开发\实验二\实验二+杨楠2024110539.docx'
shutil.copyfile(src, dst)

d = docx.Document(dst)
SHOTS = r'E:\Android exeperise\exeperise2\_shots'

def set_para_text(p, text):
    for r in list(p.runs):
        r.text = ''
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)

# 表头
set_para_text(d.paragraphs[0], '实验编号： 2  Android移动应用开发 实验报告      2026 年 10 月 10 日')
set_para_text(d.paragraphs[1], '计算机科学学院  2024 级  计科软工五班   实验名称：    UI控件代码生成以及MVC      ')
set_para_text(d.paragraphs[2], '姓名：杨楠      学号：2024110539     指导老师：李贵洋     实验成绩:_____\t')
set_para_text(d.paragraphs[5], 'Gitee网址：（待 push 后填写）')

# P28: 核心代码说明文字
code_desc = (
    '核心代码说明：\n'
    '（1）在 app/build.gradle 的 android{} 块中开启 buildFeatures { viewBinding = true }，通过 ViewBinding 绑定界面。\n'
    '（2）strings.xml 中集中定义所有用户可见文字，布局和代码均通过 @string/xxx 或 getString(R.string.xxx) 引用，未硬编码字符串。\n'
    '（3）activity_main.xml 采用垂直 LinearLayout：区一为“添加一条动态文本”按钮 + ScrollView 嵌套 LinearLayout（id=dynamicContainer）；区二为绿色背景 LinearLayout，内含“新语言推荐”提示、“确定”按钮（id=btnQuery）和结果 TextView（id=tvResult）。\n'
    '（4）ProgramAdviser.kt 作为 Model 层，内部用 Map 存储推荐语言数据，对外提供 getAdvice() 方法返回字符串资源 id。\n'
    '（5）MainActivity.kt 作为 Controller，核心代码如下图所示。'
)
p = d.paragraphs[28]
r = p.add_run(code_desc)
r.font.size = Pt(11)

# P30: MainActivity.kt 核心代码截图
p = d.paragraphs[30]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(os.path.join(SHOTS, '06_main_activity_controller.png'), width=Inches(5.8))

# P32: 运行效果说明
p = d.paragraphs[32]
r = p.add_run('图2  App 运行效果：顶部标题为“杨楠2024110539”；连续点击“添加一条动态文本”，橙色区域动态生成多条 TextView（图中已生成 6 条），内容超出后可上下滚动；点击区二“确定”按钮后，绿色区域下方以白色粗体显示推荐语言 kotlin。')
r.font.size = Pt(11)

# P33: 运行效果截图
p = d.paragraphs[33]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run()
run.add_picture(os.path.join(SHOTS, '01_running_app.png'), width=Inches(6.0))

# P46 分析
analysis = (
    '本次实验完成了题目要求的两个功能区：区一通过按钮交互，在代码中 new TextView 并 addView 到 ScrollView 嵌套的 LinearLayout 中，实现动态生成和滚动显示；区二按 MVC 模式实现交互式 ProgramAdviser，Model 层（ProgramAdviser）只负责存储推荐数据和提供 getAdvice() 查询，'
    'View 层由 activity_main.xml 承担，Controller 层（MainActivity）通过 ViewBinding 绑定视图、响应按钮点击、调用 Model 取得结果并回显。'
    '实验过程中遇到的问题及解决：①MainActivity 中 import 语句最初误用了示例包名 com.example.lab2，导致 ActivityMainBinding 报红，改为本人项目包名 cn.edu.sicnu.cs.yn.exeperise2 后通过；'
    '②初次替换布局 XML 时标签未闭合，AAPT 报“XML 文档结构必须从头至尾包含在同一个实体内”，将 activity_main.xml 整文件重写、核对所有 LinearLayout 标签闭合后解决；'
    '③app_name 在 strings.xml 中设置为“杨楠2024110539”，满足 App 标题=姓名+学号的提交要求。'
    '最终运行结果与预期一致：动态添加文本、滚动显示、按钮事件响应、Model 查询返回结果均正常，所有可见文字均走字符串资源。'
)
p = d.paragraphs[46]
r = p.add_run(analysis)
r.font.size = Pt(11)

d.save(dst)
print('saved:', dst, os.path.getsize(dst))
